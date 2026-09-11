"""Build local Markdown/DOCX from report/content.json using the official template.

Preserves the template's style definitions, page setup, title/abstract tables,
footnote and direct-format prototypes. Requires python-docx==1.2.0.
Export the DOCX to PDF with LibreOffice or Word after building.
"""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
from lxml import etree

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]


def replace_paragraph(paragraph, text, plain=False, preserve_footnote=False):
    """Replace text while retaining native paragraph and first run formatting."""
    rp = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    footnotes = [deepcopy(r._r) for r in paragraph.runs if r._r.find(qn('w:footnoteReference')) is not None] if preserve_footnote else []
    paragraph.clear()
    if plain:
        for parent in [rp, paragraph._p.pPr.find(qn('w:rPr')) if paragraph._p.pPr is not None else None]:
            if parent is not None:
                for tag in ('w:i', 'w:iCs', 'w:b', 'w:bCs', 'w:u'):
                    for element in list(parent.findall(qn(tag))):
                        parent.remove(element)
    run = paragraph.add_run(text)
    if rp is not None:
        run._r.insert(0, rp)
    for r in footnotes:
        paragraph._p.append(r)
    return paragraph


def markdown(blocks):
    lines = []
    for kind, value in blocks:
        if kind == 'title': lines.append('# ' + value)
        elif kind == 'h2': lines.append('## ' + value)
        elif kind == 'h3': lines.append('### ' + value)
        elif kind == 'abstract': lines.extend(['## Abstract', '', value])
        elif kind == 'image': lines.append(f'![Experimental results](../{value})')
        elif kind == 'table':
            lines.append('| ' + ' | '.join(value[0]) + ' |')
            lines.append('|' + '---|' * len(value[0]))
            lines.extend('| ' + ' | '.join(row) + ' |' for row in value[1:])
        else: lines.append(value)
        lines.append('')
    return '\n'.join(lines)


def build(template):
    doc = Document(template)
    blocks = [b for group in json.loads((ROOT / 'report/content.json').read_text()) for b in group]
    values = {k: t for k, t in blocks if k in ('title', 'author', 'abstract')}
    if len(values['abstract'].split()) != 150:
        raise ValueError('Submission abstract must have exactly 150 whitespace-separated words')
    # Inspect native prototypes before removing instructional body paragraphs.
    heading = deepcopy(next(p._p for p in doc.paragraphs if p.text == '1. Introduction'))
    subheading = deepcopy(next(p._p for p in doc.paragraphs if p.text == 'Limitations'))
    body = deepcopy(next(p._p for p in doc.paragraphs if p.text.startswith('What problem are you addressing')))
    header = doc.tables[0]._tbl
    title = next(Paragraph(p, doc._body) for p in header.iter(qn('w:p')) if 'PROJECT TITLE' in ''.join(p.itertext()))
    replace_paragraph(title, values['title'], preserve_footnote=True)
    # Keep title/abstract frame; collapse the six example author slots to one.
    author_tables = [t for t in header.iter(qn('w:tbl')) if not list(t.iterchildren(qn('w:tbl'))) and
                     any('Author name 1' in (x.text or '') for x in t.iter(qn('w:t')))]
    author_xml = author_tables[-1]
    authors = Table(author_xml, doc._body)
    author_cell = authors.rows[0].cells[0].merge(authors.rows[0].cells[-1])
    p = author_cell.paragraphs[0]
    replace_paragraph(p, values['author'], plain=True)
    for child in list(author_cell._tc):
        if child.tag != qn('w:tcPr') and child is not p._p:
            author_cell._tc.remove(child)
    for row in list(authors.rows)[1:]:
        author_xml.remove(row._tr)
    abstract_p = next(Paragraph(p, doc._body) for p in header.iter(qn('w:p'))
                      if any('Summarize your project' in (t.text or '') for t in p.iter(qn('w:t'))))
    replace_paragraph(abstract_p, values['abstract'], plain=True)
    for child in list(doc._element.body):
        if child is not header and child.tag != qn('w:sectPr'):
            doc._element.body.remove(child)

    def para(text, prototype=body, plain=True):
        node = deepcopy(prototype)
        doc._element.body.insert_element_before(node, 'w:sectPr')
        return replace_paragraph(Paragraph(node, doc._body), text, plain=plain)

    for kind, value in blocks:
        if kind in ('title', 'author', 'abstract'): continue
        if kind in ('h2', 'h3'):
            p = para(value, heading if kind == 'h2' else subheading, False)
            p.paragraph_format.keep_with_next = True
            if value in ('References', 'Appendix: Limitations and Dual-Use Considerations'):
                p.paragraph_format.page_break_before = True
        elif kind == 'image':
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            section = doc.sections[0]
            p.add_run().add_picture(str(ROOT / value), width=section.page_width-section.left_margin-section.right_margin)
        elif kind == 'table':
            table = doc.add_table(rows=0, cols=len(value[0]))
            table.autofit = True
            for i, row in enumerate(value):
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    p = cell.paragraphs[0]
                    p._p.getparent().replace(p._p, deepcopy(body))
                    p = replace_paragraph(cell.paragraphs[0], text, plain=True)
                    p.paragraph_format.keep_with_next = True
                    if i == 0:
                        for run in p.runs: run.bold = True
                if i == 0:
                    prop = OxmlElement('w:tblHeader')
                    table.rows[-1]._tr.get_or_add_trPr().append(prop)
        else:
            p = para(value)
            if kind == 'caption':
                for run in p.runs: run.italic = True
    doc.core_properties.title = values['title']
    doc.core_properties.author = 'Matías Podeley — BAISH'
    doc.core_properties.subject = 'AI Incident Response Sprint, Track 1 — Containment'
    target = ROOT / 'report/agent-delegate.docx'
    doc.save(target)
    (ROOT / 'report/report.md').write_text(markdown(blocks))
    before = ZipFile(template).read('word/styles.xml')
    after = ZipFile(target).read('word/styles.xml')
    if etree.tostring(etree.fromstring(before)) != etree.tostring(etree.fromstring(after)):
        raise ValueError('Template style definitions changed')
    # python-docx rewrites the XML declaration only; keep the original bytes too.
    original = target.read_bytes()
    with ZipFile(BytesIO(original)) as source, ZipFile(target, 'w') as destination:
        for info in source.infolist():
            destination.writestr(info, before if info.filename == 'word/styles.xml' else source.read(info.filename))
    after = ZipFile(target).read('word/styles.xml')
    proof = {'template_styles_sha256': hashlib.sha256(before).hexdigest(),
             'output_styles_sha256': hashlib.sha256(after).hexdigest(),
             'styles_identical': before == after,
             'section_properties_identical': doc.sections[0]._sectPr.xml == Document(template).sections[0]._sectPr.xml,
             'abstract_words': len(values['abstract'].split()),
             'adaptations': ['One author instead of six placeholders', 'Instructional body replaced',
                             'Required Limitations and Dual-Use Considerations appendix',
                             'Added results tables and figure; retained native heading/body formatting'],
             'template_url': json.loads((ROOT / 'report/template-provenance.json').read_text())['url']}
    (ROOT / 'report/template-check.json').write_text(json.dumps(proof, indent=2) + '\n')
    print(target)


if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--template', type=Path, required=True)
    build(p.parse_args().template)
